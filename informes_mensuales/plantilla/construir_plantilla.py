"""Construye plantilla/informe_esmt.docx a partir de un informe manual ya aprobado.

Conserva lo que diseñó el equipo (portada, header/footer, estilos, colores ESMT,
índice) y sustituye el contenido del cliente por campos docxtpl ({{ }}, {%p %}, {%tr %}).
Se ejecuta una sola vez; después la plantilla se puede retocar directamente en Word
(textos fijos, colores, tamaños) siempre que no se rompan las etiquetas {{ ... }}.

Uso:
    .venv/bin/python plantilla/construir_plantilla.py "/ruta/INFORME MENSUAL SEPTIEMBRE STEREN.docx"
"""
import re
import sys
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

DESTINO = Path(__file__).resolve().parent / "informe_esmt.docx"
ANCHO = 8500  # ancho útil de la página en twips (A4 con márgenes de 3 cm)

NARANJA, NARANJA_TITULO, CREMA, BORDE, TEXTO, OSCURO = "EA580C", "FF4F1F", "FFF7ED", "FED7AA", "333333", "0F172A"
W = nsdecls("w")


# ---------------------------------------------------------------- bloques XML
def _run(texto, sz=21, color=TEXTO, negrita=False, cursiva=False):
    extra = ("<w:b/>" if negrita else "") + ("<w:i/>" if cursiva else "")
    return (f'<w:r><w:rPr><w:rFonts w:ascii="Calibri" w:hAnsi="Calibri"/>{extra}<w:color w:val="{color}"/>'
            f'<w:sz w:val="{sz}"/><w:lang w:val="es-MX"/></w:rPr><w:t xml:space="preserve">{escape(texto)}</w:t></w:r>')


def parrafo(texto, negrita=False, centrado=False, sz=21, color=TEXTO, cursiva=False, antes=40, despues=120, sangria=None):
    jc = '<w:jc w:val="center"/>' if centrado else '<w:jc w:val="both"/>'
    ind = f'<w:ind w:left="{sangria[0]}" w:hanging="{sangria[1]}"/>' if sangria else ""
    return parse_xml(f'<w:p {W}><w:pPr><w:spacing w:before="{antes}" w:after="{despues}"/>{ind}{jc}</w:pPr>'
                     f'{_run(texto, sz, color, negrita, cursiva)}</w:p>')


def control(etiqueta):
    """Párrafo que solo contiene una etiqueta {%p ... %}; docxtpl lo elimina al renderizar."""
    return parse_xml(f'<w:p {W}><w:r><w:t xml:space="preserve">{escape(etiqueta)}</w:t></w:r></w:p>')


def titulo1(texto):
    rpr = ('<w:rFonts w:ascii="Calibri" w:eastAsia="Times New Roman" w:hAnsi="Calibri" w:cs="Times New Roman"/>'
           '<w:bCs w:val="0"/><w:color w:val="FFFFFF"/><w:kern w:val="20"/><w:sz w:val="24"/><w:szCs w:val="24"/><w:lang w:val="es-MX"/>')
    return parse_xml(
        f'<w:p {W}><w:pPr><w:pStyle w:val="Heading1"/><w:numPr><w:ilvl w:val="0"/><w:numId w:val="1"/></w:numPr>'
        '<w:pBdr><w:top w:val="single" w:sz="48" w:space="0" w:color="FFFFFF"/><w:left w:val="single" w:sz="6" w:space="0" w:color="FFFFFF"/>'
        '<w:bottom w:val="single" w:sz="6" w:space="2" w:color="FFFFFF"/></w:pBdr>'
        f'<w:shd w:val="clear" w:color="auto" w:fill="{NARANJA_TITULO}"/><w:spacing w:before="360" w:after="120" w:line="240" w:lineRule="atLeast"/>'
        f'<w:ind w:hanging="720"/><w:jc w:val="both"/><w:rPr>{rpr}</w:rPr></w:pPr>'
        f'<w:r><w:rPr>{rpr}</w:rPr><w:t xml:space="preserve">{escape(texto)}</w:t></w:r></w:p>')


def titulo2(texto):
    return parse_xml(f'<w:p {W}><w:pPr><w:pStyle w:val="Heading2"/></w:pPr>'
                     f'<w:r><w:rPr><w:rFonts w:ascii="Calibri" w:hAnsi="Calibri"/><w:color w:val="{NARANJA}"/><w:sz w:val="22"/></w:rPr>'
                     f'<w:t xml:space="preserve">{escape(texto)}</w:t></w:r></w:p>')


def imagen(variable):
    return parrafo("{{ " + variable + " }}", centrado=True, antes=120, despues=60)


def fuente(texto):
    return parrafo(texto, centrado=True, sz=17, color="7F7F7F", cursiva=True, antes=0, despues=200)


def salto_pagina():
    return parse_xml(f'<w:p {W}><w:r><w:br w:type="page"/></w:r></w:p>')


def _celda(texto, ancho, alin, tipo, sz=None):
    fondo = {"enc": NARANJA, "total": CREMA}.get(tipo)
    shd = f'<w:shd w:val="clear" w:color="auto" w:fill="{fondo}"/>' if fondo else ""
    m = {"enc": 120, "total": 100}.get(tipo, 90)
    color = {"enc": "FFFFFF", "total": OSCURO}.get(tipo, TEXTO)
    tam = sz or (20 if tipo in ("enc", "total") else 19)
    jc = {"c": "center", "d": "right"}.get(alin)
    ppr = f'<w:pPr><w:jc w:val="{jc}"/></w:pPr>' if jc else ""
    return (f'<w:tc><w:tcPr><w:tcW w:w="{ancho}" w:type="dxa"/>{shd}<w:tcMar><w:top w:w="{m}" w:type="dxa"/>'
            f'<w:left w:w="120" w:type="dxa"/><w:bottom w:w="{m}" w:type="dxa"/><w:right w:w="120" w:type="dxa"/></w:tcMar></w:tcPr>'
            f'<w:p>{ppr}{_run(texto, tam, color, negrita=tipo in ("enc", "total"))}</w:p></w:tc>')


def _fila(celdas, encabezado=False):
    trpr = '<w:trPr><w:cantSplit/>' + ('<w:tblHeader/>' if encabezado else '') + '<w:jc w:val="center"/></w:trPr>'
    return f'<w:tr>{trpr}{"".join(celdas)}</w:tr>'


def _fila_control(etiqueta, columnas):
    return _fila([_celda(etiqueta if i == 0 else "", ancho, "", "dato") for i, (_, ancho, _, _) in enumerate(columnas)])


def tabla(columnas, bucle, total=None, sz=None):
    """columnas: [(encabezado, ancho, alineación 'i'|'c'|'d', expresión jinja de la celda)]"""
    grid = "".join(f'<w:gridCol w:w="{c[1]}"/>' for c in columnas)
    filas = [_fila([_celda(enc, ancho, alin, "enc", sz) for enc, ancho, alin, _ in columnas], encabezado=True),
             _fila_control("{%tr for " + bucle + " %}", columnas),
             _fila([_celda("{{ " + expr + " }}", ancho, alin, "dato", sz) for _, ancho, alin, expr in columnas]),
             _fila_control("{%tr endfor %}", columnas)]
    if total:
        filas.append(_fila([_celda(t, c[1], c[2], "total", sz) for t, c in zip(total, columnas)]))
    return parse_xml(
        f'<w:tbl {W}><w:tblPr><w:tblW w:w="0" w:type="auto"/><w:jc w:val="center"/>'
        f'<w:tblBorders><w:top w:val="single" w:sz="4" w:space="0" w:color="{BORDE}"/>'
        f'<w:bottom w:val="single" w:sz="8" w:space="0" w:color="{NARANJA}"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{BORDE}"/></w:tblBorders>'
        '<w:tblLayout w:type="fixed"/><w:tblLook w:val="04A0" w:firstRow="1" w:lastRow="0" w:firstColumn="1" w:lastColumn="0" w:noHBand="0" w:noVBand="1"/>'
        f'</w:tblPr><w:tblGrid>{grid}</w:tblGrid>{"".join(filas)}</w:tbl>')


# ------------------------------------------------------------- contenido
def cuerpo():
    b = []
    a = b.append

    # 1. Resumen de soporte
    a(titulo1("RESUMEN DE SOPORTE"))
    a(control("{%p for p in textos.soporte.intro %}")); a(parrafo("{{ p }}")); a(control("{%p endfor %}"))
    a(control("{%p for t in soporte.tickets %}"))
    a(titulo2("{{ t.asunto }} (#{{ t.id }})"))
    a(parrafo("Estado: {{ t.etapa }}{% if t.en_curso %} · en curso{% endif %} · Tipo: {{ t.tipo }}", sz=18, color="7F7F7F", antes=0, despues=60))
    a(parrafo("{{ t.resumen }}", antes=0))
    a(control("{%p endfor %}"))

    # 2. Volumen de casos
    a(titulo1("VOLUMEN DE CASOS"))
    a(parrafo("Se comparte la cantidad de casos para este cliente de acuerdo con los tipos de casos abiertos en el mes de {{ mes }} de {{ anio }}."))
    a(tabla([("Tipo de caso", 6000, "i", "r.tipo"), ("Cantidad", 2500, "c", "r.cantidad")],
            "r in soporte.por_tipo", total=["Total", "{{ soporte.total }}"]))
    a(imagen("img.tipos"))

    a(titulo2("TIEMPO PROMEDIO"))
    a(parrafo("{{ soporte.texto_promedio }}"))
    a(tabla([("ID", 800, "c", "t.id"), ("Asunto", 3900, "i", "t.asunto"), ("Creado el", 1250, "c", "t.creado_fmt"),
             ("Cierre", 1600, "c", "t.cierre_fmt"), ("Días", 950, "c", "t.dias_fmt")],
            "t in soporte.tickets", total=["", "", "", "Días promedio", "{{ soporte.promedio_fmt }}"], sz=18))
    a(parrafo("{{ soporte.nota_promedio }}", sz=17, color="7F7F7F", cursiva=True, antes=60))

    a(titulo2("HISTORIAL DE HORAS"))
    a(parrafo("Horas registradas en los casos aperturados en {{ mes }} de {{ anio }}."))
    a(tabla([("ID", 800, "c", "t.id"), ("Asunto", 4400, "i", "t.asunto"), ("Estado", 2100, "c", "t.etapa"), ("Horas", 1200, "c", "t.horas_fmt")],
            "t in soporte.tickets", total=["", "", "Total de horas", "{{ soporte.horas_fmt }}"], sz=18))

    a(titulo2("ESTADO DE LOS TICKETS"))
    a(parrafo("{{ soporte.texto_estados }}"))
    a(tabla([("Estado", 6000, "i", "r.etapa"), ("Cantidad", 2500, "c", "r.cantidad")],
            "r in soporte.por_etapa", total=["Total", "{{ soporte.total }}"]))
    a(imagen("img.estados"))

    # 3. Licenciamiento (se omite entera si el informe se genera sin License Manager)
    a(control("{%p if not lic_omitido %}"))
    a(titulo1("LICENCIAMIENTO (REPORTE DE USO)"))
    a(control("{%p if lic %}"))
    a(parrafo("Organización: {{ lic.organizacion }} · Periodo de facturación: {{ lic.periodo_txt }}", negrita=True))
    a(titulo2("Resumen ejecutivo de uso"))
    a(control("{%p for p in textos.licenciamiento.resumen %}")); a(parrafo("{{ p }}")); a(control("{%p endfor %}"))

    a(titulo2("Estado de KPIs principales"))
    a(parrafo("Licencias", negrita=True, antes=120))
    a(tabla([("Licencia", 3700, "i", "l.nombre"), ("Contratadas", 1200, "c", "l.contratadas_fmt"), ("En uso", 1100, "c", "l.en_uso_fmt"),
             ("Sobreuso", 1200, "c", "l.sobreuso_fmt"), ("% de uso", 1300, "c", "l.pct_fmt")],
            "l in lic.licencias", sz=18))
    a(parrafo("Recursos y consumo", negrita=True, antes=200))
    a(tabla([("Recurso", 3300, "i", "r.nombre"), ("Disponible", 1500, "c", "r.disponible_fmt"), ("En uso", 1400, "c", "r.en_uso_fmt"),
             ("% de uso", 1000, "c", "r.pct_fmt"), ("Estado", 1300, "c", "r.estado")],
            "r in lic.recursos", sz=18))
    a(parrafo("{{ lic.nota_recursos }}", sz=17, color="7F7F7F", cursiva=True, antes=60))

    a(titulo2("Alertas"))
    a(control("{%p for x in lic.alertas %}"))
    a(parrafo("•  {{ x.tema }}: {{ x.detalle }}", sangria=(360, 200), despues=60))
    a(control("{%p endfor %}"))
    a(control("{%p for p in textos.licenciamiento.alertas %}")); a(parrafo("{{ p }}")); a(control("{%p endfor %}"))

    a(titulo2("CAPACIDAD DISPONIBLE VS UTILIZADA"))
    a(imagen("img.capacidad"))
    a(fuente("Fuente: License Manager ESMT (API de facturación de Genesys Cloud)."))

    a(control("{%p if lic.historico %}"))
    a(titulo2("HISTÓRICO DE CONSUMO {{ anio }}"))
    a(parrafo("Uso por periodo de facturación y porcentaje sobre lo disponible en cada periodo."))
    a(tabla([("Periodo", 1300, "i", "h.periodo"), ("{{ lic.historico_licencia }}", 1100, "c", "h.licencia"), ("BYOC (min)", 1150, "c", "h.byoc"),
             ("IVR (min)", 1050, "c", "h.ivr"), ("API", 1150, "c", "h.api"), ("Storage", 900, "c", "h.storage"),
             ("IA Tokens", 950, "c", "h.ia_tokens"), ("Transcr.", 900, "c", "h.transcripcion")],
            "h in lic.historico", sz=15))
    a(control("{%p endif %}"))
    a(control("{%p else %}"))
    a(parrafo("La información de licenciamiento no estaba disponible al generar este borrador."))
    a(control("{%p endif %}"))
    a(control("{%p endif %}"))

    # 4. Casos mes a mes
    a(titulo1("CASOS MES A MES"))
    a(parrafo("Cantidad de casos mes a mes diferenciando por tipo de caso. El comportamiento en {{ anio }} ha sido el siguiente:"))
    a(tabla([("Mes", 2300, "i", "m.mes"), ("Incidente", 1500, "c", "m.incidente"), ("Requerimiento", 1700, "c", "m.requerimiento"),
             ("Cambio", 1400, "c", "m.cambio"), ("Total", 1500, "c", "m.total")],
            "m in serie", total=["Total", "{{ serie_total.incidente }}", "{{ serie_total.requerimiento }}", "{{ serie_total.cambio }}", "{{ serie_total.total }}"]))
    a(imagen("img.mes_a_mes"))

    # 5. Disponibilidad
    a(titulo1("DISPONIBILIDAD GENERAL DEL SERVICIO"))
    a(control("{%p if estado %}"))
    a(control("{%p for p in estado.parrafos %}")); a(parrafo("{{ p }}")); a(control("{%p endfor %}"))
    a(titulo2("Disponibilidad de la plataforma en {{ anio }}"))
    a(tabla([("Mes", 3000, "i", "d.mes"), ("{{ estado.region }}", 3000, "c", "d.valor")], "d in estado.disponibilidad"))
    a(titulo2("Incidentes publicados en la región"))
    a(control("{%p if estado.incidentes %}"))
    a(tabla([("Inicio (UTC)", 1700, "c", "i.fecha"), ("Incidente", 3900, "i", "i.titulo"), ("Impacto", 1200, "c", "i.impacto"),
             ("Duración", 1700, "c", "i.duracion")], "i in estado.incidentes", sz=18))
    a(control("{%p else %}"))
    a(parrafo("No se publicaron incidentes para la región durante el mes."))
    a(control("{%p endif %}"))
    a(fuente("Fuente: https://status.mypurecloud.com/ (consultado el {{ fecha_corte }}). La disponibilidad de la plataforma puede no reflejar la disponibilidad de cada cliente."))
    a(control("{%p else %}"))
    a(parrafo("La información de disponibilidad no estaba disponible al generar este borrador."))
    a(control("{%p endif %}"))
    return b


def indice():
    """Índice prellenado (sin números de página); lo regeneran LibreOffice al exportar a PDF y Word al abrir."""
    return [
        control("{%p for e in indice %}"),
        control("{%p if e.nivel == 1 %}"),
        parse_xml(f'<w:p {W}><w:pPr><w:pStyle w:val="TOC1"/></w:pPr>{_run("{{ e.texto }}", 22, OSCURO, negrita=True)}</w:p>'),
        control("{%p else %}"),
        parse_xml(f'<w:p {W}><w:pPr><w:pStyle w:val="TOC2"/></w:pPr>{_run("{{ e.texto }}", 20, TEXTO)}</w:p>'),
        control("{%p endif %}"),
        control("{%p endfor %}"),
    ]


# ------------------------------------------------------------- construcción
def congelar_includepicture(raiz):
    """Deja fija la imagen ya incrustada y elimina el campo INCLUDEPICTURE que la envuelve.

    Los informes manuales traen en el header campos INCLUDEPICTURE (anidados) con URLs de
    miniaturas de Google y otros sitios. Si Word actualiza campos al abrir, descarga esas
    imágenes a tamaño completo, el header crece hasta ocupar la página y el documento
    pasa de ~11 a cientos de páginas.
    """
    pila, quitar, campos = [], [], 0
    for r in list(raiz.iter(qn("w:r"))):
        fc, instr = r.find(qn("w:fldChar")), r.find(qn("w:instrText"))
        if fc is not None:
            tipo = fc.get(qn("w:fldCharType"))
            if tipo == "begin":
                pila.append({"instr": "", "runs": [r]})
            elif pila:
                pila[-1]["runs"].append(r)
                if tipo == "end":
                    campo = pila.pop()
                    if "INCLUDEPICTURE" in campo["instr"].upper():
                        quitar.extend(campo["runs"])
                        campos += 1
        elif instr is not None and pila:
            pila[-1]["instr"] += instr.text or ""
            pila[-1]["runs"].append(r)
    for r in quitar:
        r.getparent().remove(r)
    for simple in list(raiz.iter(qn("w:fldSimple"))):
        if "INCLUDEPICTURE" in (simple.get(qn("w:instr")) or "").upper():
            for hijo in list(simple):
                simple.addprevious(hijo)
            simple.getparent().remove(simple)
            campos += 1
    return campos


def reemplazar_texto(parrafos, viejo, nuevo):
    for p in parrafos:
        for r in p.runs:
            if r.text.strip().lower() == viejo.lower():
                r.text = r.text.replace(r.text.strip(), nuevo)


def igualar_word_y_libreoffice(doc):
    """Retoques que no cambian lo que muestra Word y hacen que el PDF de LibreOffice se vea igual.

    - LibreOffice regenera el título del índice con el estilo «TOC Heading»: ese estilo pasa a tener
      el formato del título (antes «Intense Quote» + formato directo).
    - Las entradas de nivel 2 del índice: Word copia la fuente Calibri de los títulos; LibreOffice usa
      la del estilo TOC2 (Times, heredada de Normal), así que TOC2 declara Calibri.
    - El header trae un «PROYECTO» que en Word queda tapado por el logo y LibreOffice dibuja encima.
    """
    estilos = doc.styles.element
    nuevo = parse_xml(
        f'<w:style {W} w:type="paragraph" w:styleId="TOCHeading"><w:name w:val="TOC Heading"/>'
        '<w:basedOn w:val="Normal"/><w:next w:val="Normal"/><w:uiPriority w:val="39"/><w:unhideWhenUsed/><w:qFormat/>'
        '<w:pPr><w:keepNext/><w:pageBreakBefore/><w:pBdr><w:bottom w:val="single" w:sz="4" w:space="4" w:color="4472C4" '
        'w:themeColor="accent1"/></w:pBdr><w:spacing w:before="200" w:after="280" w:line="276" w:lineRule="auto"/>'
        '<w:ind w:left="0" w:right="936"/><w:outlineLvl w:val="9"/></w:pPr>'
        '<w:rPr><w:rFonts w:ascii="Calibri" w:hAnsi="Calibri" w:cs="Calibri"/><w:b/><w:bCs/><w:color w:val="F85E00"/>'
        '<w:sz w:val="48"/><w:szCs w:val="48"/><w:lang w:val="es-ES"/></w:rPr></w:style>')
    viejo = next((e for e in estilos.iter(qn("w:style")) if e.get(qn("w:styleId")) == "TOCHeading"), None)
    if viejo is not None:
        viejo.addprevious(nuevo)
        estilos.remove(viejo)
    else:
        estilos.append(nuevo)
    toc2 = next(e for e in estilos.iter(qn("w:style")) if e.get(qn("w:styleId")) == "TOC2")
    rpr = toc2.find(qn("w:rPr"))
    if rpr is None:
        rpr = parse_xml(f"<w:rPr {W}/>")
        toc2.append(rpr)
    if rpr.find(qn("w:rFonts")) is None:
        rpr.insert(0, parse_xml(f'<w:rFonts {W} w:ascii="Calibri" w:hAnsi="Calibri" w:cs="Calibri"/>'))

    sdt = next(doc.element.body.iter(qn("w:sdt")))
    titulo = next(sdt.iter(qn("w:p")))
    titulo.find(qn("w:pPr")).find(qn("w:pStyle")).set(qn("w:val"), "TOCHeading")

    for s in doc.sections:
        for parte in (s.header, s.first_page_header, s.even_page_header):
            reemplazar_texto(parte.paragraphs, "PROYECTO", "")


def construir(origen):
    doc = Document(origen)
    body = doc.element.body
    hijos = list(body)
    i_sdt = next(i for i, e in enumerate(hijos) if e.tag == qn("w:sdt"))

    # Portada
    portada = [p for p in doc.paragraphs if p._p in hijos[:i_sdt]]
    reemplazar_texto(portada, "SEPTIEMBRE", "{{ mes_mayus }} {{ anio }}")
    reemplazar_texto(portada, "Support Team", "{{ preparado_por }}")
    reemplazar_texto(portada, "Grupo ESMT", "{{ empresa }}")
    logo_original = None
    for p in portada:
        if any("INCLUDEPICTURE" in (t.text or "") for t in p._p.iter(qn("w:instrText"))):
            blip = next(p._p.iter(qn("a:blip")), None)
            if blip is not None:
                logo_original = doc.part.related_parts[blip.get(qn("r:embed"))]
            p._p.addprevious(imagen("logo_cliente"))
            body.remove(p._p)

    # Header y footer: el mes y el logo del cliente (se marca para que render.py lo reemplace)
    for s in doc.sections:
        for parte in (s.header, s.first_page_header, s.even_page_header, s.footer, s.first_page_footer, s.even_page_footer):
            reemplazar_texto(parte.paragraphs, "Septiembre", "{{ mes }}")
            for blip in parte._element.iter(qn("a:blip")):
                if logo_original is not None and parte.part.related_parts.get(blip.get(qn("r:embed"))) is logo_original:
                    grafico = blip.getparent().getparent().getparent()  # pic:pic → a:graphicData
                    for nodo in grafico.iter(qn("pic:cNvPr")):
                        nodo.set("descr", "logo_cliente")
                    dibujo = next(a for a in blip.iterancestors() if a.tag in (qn("wp:inline"), qn("wp:anchor")))
                    dibujo.find(qn("wp:docPr")).set("descr", "logo_cliente")

    # Índice: se conserva el título y el campo TOC; las entradas se rellenan al renderizar
    sdt = hijos[i_sdt]
    contenido = sdt.find(qn("w:sdtContent"))
    pars = contenido.findall(qn("w:p"))
    instr = next(t.text for t in sdt.iter(qn("w:instrText")) if "TOC" in (t.text or ""))
    for p in pars[1:]:
        contenido.remove(p)
    # El índice empieza en página nueva, sin depender del tamaño del logo de la portada
    ppr = pars[0].find(qn("w:pPr"))
    if ppr is not None and ppr.find(qn("w:pageBreakBefore")) is None:
        ppr.insert(1, parse_xml(f'<w:pageBreakBefore {W}/>'))
    inicio = parse_xml(f'<w:p {W}><w:pPr><w:pStyle w:val="TOC1"/><w:spacing w:after="0"/></w:pPr>'
                       f'<w:r><w:fldChar w:fldCharType="begin" w:dirty="true"/></w:r>'
                       f'<w:r><w:instrText xml:space="preserve">{escape(instr)}</w:instrText></w:r>'
                       f'<w:r><w:fldChar w:fldCharType="separate"/></w:r></w:p>')
    fin = parse_xml(f'<w:p {W}><w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>')
    for e in [inicio, *indice(), fin]:
        contenido.append(e)

    # Cuerpo: se elimina el contenido del cliente y se inserta la estructura con campos
    sect = body.find(qn("w:sectPr"))
    for e in list(body)[list(body).index(sdt) + 1:]:
        if e is not sect:
            body.remove(e)
    for e in [salto_pagina(), *cuerpo()]:
        sect.addprevious(e)

    # Campos que descargan imágenes de internet → imágenes fijas (cuerpo, headers y footers)
    partes = [doc.element] + [rel.target_part.element for rel in doc.part.rels.values()
                              if rel.reltype.endswith(("/header", "/footer"))]
    congelados = sum(congelar_includepicture(p) for p in partes)
    print(f"Campos INCLUDEPICTURE convertidos en imagen fija: {congelados}")

    # Relaciones huérfanas (imágenes, gráficas e hipervínculos del informe original)
    usados = set(re.findall(r'r:(?:embed|id|link)="(rId\d+)"', doc.element.xml))
    for rid, rel in list(doc.part.rels.items()):
        if rid not in usados and any(t in rel.reltype for t in ("/image", "/chart", "/hyperlink")):
            del doc.part.rels[rid]

    # Pedir a Word que actualice el índice al abrir
    settings = doc.settings.element
    if settings.find(qn("w:updateFields")) is None:
        settings.append(parse_xml(f'<w:updateFields {W} w:val="true"/>'))

    igualar_word_y_libreoffice(doc)

    doc.core_properties.title = "Informe mensual"
    doc.core_properties.author = "Grupo ESMT"
    doc.save(DESTINO)
    print(f"Plantilla generada: {DESTINO}")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        raise SystemExit(__doc__)
    construir(sys.argv[1])
